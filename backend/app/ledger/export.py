"""CSV and Excel export of the ledger. Cells are protected against spreadsheet formula injection."""

import csv
import io
import sqlite3
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HEADERS = [
    "Date", "Vendor", "Vendor GSTIN", "Bill no", "Category",
    "Taxable (INR)", "CGST (INR)", "SGST (INR)", "IGST (INR)", "Total (INR)",
]  # fmt: skip
_MONEY_FIELDS = ["taxable_paise", "cgst_paise", "sgst_paise", "igst_paise", "total_paise"]
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


def safe_text(value: Any) -> str:
    """Neutralise text that Excel/Sheets would run as a formula."""
    text = "" if value is None else str(value)
    return "'" + text if text.startswith(_FORMULA_STARTS) else text


def rupees(paise: int | None) -> float:
    return round((paise or 0) / 100, 2)


def _row(r: sqlite3.Row) -> list[Any]:
    return [
        r["bill_date"] or "",
        safe_text(r["vendor_name"]),
        safe_text(r["vendor_gstin"]),
        safe_text(r["bill_no"]),
        safe_text(r["category"]),
        *[rupees(r[f]) for f in _MONEY_FIELDS],
    ]


def _totals(rows: list[sqlite3.Row]) -> list[float]:
    return [round(sum((r[f] or 0) for r in rows) / 100, 2) for f in _MONEY_FIELDS]


def to_csv(rows: list[sqlite3.Row]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(HEADERS)
    for r in rows:
        writer.writerow([f"{v:.2f}" if isinstance(v, float) else v for v in _row(r)])
    writer.writerow(["TOTAL", "", "", "", "", *[f"{t:.2f}" for t in _totals(rows)]])
    return b"\xef\xbb\xbf" + buf.getvalue().encode()  # BOM so Excel reads UTF-8 correctly


def to_xlsx(rows: list[sqlite3.Row], month: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = month
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(_row(r))
    ws.append(["TOTAL", "", "", "", "", *_totals(rows)])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=10):
        for cell in row:
            cell.number_format = "#,##0.00"
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # never interpret as a formula
    widths = [12, 28, 18, 14, 18, 14, 12, 12, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
