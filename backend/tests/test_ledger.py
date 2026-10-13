import csv
import io
from datetime import date

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.ledger import service
from tests.conftest import signup_payload
from tests.helpers import add_bill, add_invoice, add_party


def setup(client: TestClient, email: str = "a@example.com") -> tuple[int, object]:
    client.post("/api/auth/signup", json=signup_payload(email))
    return client.get("/api/auth/me").json()["business_id"], client.app.state.db  # type: ignore[attr-defined]


def test_filters_search_and_pagination(client: TestClient) -> None:
    biz, conn = setup(client)
    vendor = add_party(conn, biz, "Kumar Electricals")  # type: ignore[arg-type]
    add_bill(conn, biz, date="2026-05-12", category="stock purchase")  # type: ignore[arg-type]
    add_bill(
        conn,
        biz,
        date="2026-05-20",
        category="repairs",
        vendor="Kumar Electricals",
        party_id=vendor,
        bill_no="KE-9",
    )  # type: ignore[arg-type]
    add_bill(conn, biz, date="2026-06-02", category="rent", vendor="Landlord 100%")  # type: ignore[arg-type]
    add_bill(conn, biz, status="needs_review")  # type: ignore[arg-type]
    all_ = client.get("/api/ledger").json()
    assert all_["total_count"] == 3  # confirmed only by default
    assert [e["bill_date"] for e in all_["entries"]] == ["2026-06-02", "2026-05-20", "2026-05-12"]
    assert client.get("/api/ledger?month=2026-05").json()["total_count"] == 2
    assert client.get("/api/ledger?category=repairs").json()["total_count"] == 1
    assert client.get(f"/api/ledger?vendor_id={vendor}").json()["entries"][0]["bill_no"] == "KE-9"
    assert client.get("/api/ledger?status=needs_review").json()["total_count"] == 1
    assert client.get("/api/ledger?q=kumar").json()["total_count"] == 1
    assert client.get("/api/ledger?q=KE-9").json()["total_count"] == 1
    assert client.get("/api/ledger?q=%25").json()["total_count"] == 1  # literal %, not a wildcard
    assert client.get("/api/ledger?q=_").json()["total_count"] == 0
    page = client.get("/api/ledger?limit=2&offset=2").json()
    assert page["total_count"] == 3 and len(page["entries"]) == 1
    assert client.get("/api/ledger?month=2026-13").status_code == 422
    assert client.get("/api/ledger?limit=0").status_code == 422


def test_totals_cover_all_pages(client: TestClient) -> None:
    biz, conn = setup(client)
    for _ in range(3):
        add_bill(conn, biz)  # type: ignore[arg-type]
    r = client.get("/api/ledger?limit=1").json()
    assert r["totals"]["total"] == 3 * 118000 and r["totals"]["cgst"] == 3 * 9000


def test_filter_options(client: TestClient) -> None:
    biz, conn = setup(client)
    add_party(conn, biz, "Zed Vendor")  # type: ignore[arg-type]
    add_party(conn, biz, "Some Customer", kind="customer")  # type: ignore[arg-type]
    add_bill(conn, biz, date="2026-04-01")  # type: ignore[arg-type]
    add_bill(conn, biz, date="2026-05-01")  # type: ignore[arg-type]
    o = client.get("/api/ledger/filters").json()
    assert o["months"] == ["2026-05", "2026-04"]
    assert [v["name"] for v in o["vendors"]] == ["Zed Vendor"]
    assert "rent" in o["categories"]


def test_dashboard_numbers(client: TestClient) -> None:
    biz, conn = setup(client)
    today = date(2026, 6, 15)
    add_bill(
        conn, biz, date="2026-06-02", category="rent", taxable=200000, cgst=0, sgst=0, igst=36000
    )  # type: ignore[arg-type]
    add_bill(conn, biz, date="2026-06-10", category="rent")  # type: ignore[arg-type]
    add_bill(conn, biz, date="2026-06-11", category="repairs")  # type: ignore[arg-type]
    add_bill(conn, biz, date="2026-05-31", category="repairs")  # last month: excluded  # type: ignore[arg-type]
    add_bill(conn, biz, date="2026-06-12", status="needs_review")  # type: ignore[arg-type]
    cust = add_party(conn, biz, "Cust", kind="customer")  # type: ignore[arg-type]
    add_invoice(conn, biz, cust, due_date="2026-06-14", total=100000, paid=40000)  # type: ignore[arg-type]
    add_invoice(conn, biz, cust, invoice_no="S-2", due_date="2026-06-15")  # due today: not overdue  # type: ignore[arg-type]
    add_invoice(
        conn, biz, cust, invoice_no="S-3", due_date="2026-01-01", status="paid", paid=500000
    )  # type: ignore[arg-type]
    d = service.dashboard(conn, biz, today)  # type: ignore[arg-type]
    assert d["month"] == "2026-06" and d["bills"] == 3
    assert d["purchases_paise"] == 236000 + 118000 + 118000
    assert d["input_gst"] == {
        "cgst_paise": 18000,
        "sgst_paise": 18000,
        "igst_paise": 36000,
        "total_paise": 72000,
    }
    assert d["top_categories"][0] == {"category": "rent", "total_paise": 354000}
    assert d["needs_review"] == 1
    assert d["overdue"] == {"count": 1, "amount_paise": 60000}
    assert client.get("/api/dashboard").status_code == 200


def test_dashboard_empty(client: TestClient) -> None:
    setup(client)
    d = client.get("/api/dashboard").json()
    assert d["purchases_paise"] == 0 and d["top_categories"] == [] and d["overdue"]["count"] == 0


def test_csv_export(client: TestClient) -> None:
    biz, conn = setup(client)
    add_bill(conn, biz, date="2026-05-12", vendor='=HYPERLINK("http://evil","x")')  # type: ignore[arg-type]
    add_bill(conn, biz, date="2026-05-20", vendor="Plain, Vendor", bill_no='A"B')  # type: ignore[arg-type]
    add_bill(conn, biz, date="2026-06-01")  # other month  # type: ignore[arg-type]
    r = client.get("/api/ledger/export?month=2026-05&format=csv")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert "kanakku-ledger-2026-05.csv" in r.headers["content-disposition"]
    assert r.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][:2] == ["Date", "Vendor"] and "CGST (INR)" in rows[0] and "IGST (INR)" in rows[0]
    assert len(rows) == 4  # header + 2 bills + total
    assert rows[1][1].startswith("'=")  # formula neutralised
    assert rows[2][1] == "Plain, Vendor" and rows[2][3] == 'A"B'
    assert rows[-1][0] == "TOTAL" and rows[-1][5] == "2000.00" and rows[-1][9] == "2360.00"


def test_xlsx_export_opens_with_numbers(client: TestClient) -> None:
    biz, conn = setup(client)
    add_bill(conn, biz, date="2026-05-12", vendor="=1+1")  # type: ignore[arg-type]
    add_bill(conn, biz, date="2026-05-20", taxable=250050, cgst=0, sgst=0, igst=45009)  # type: ignore[arg-type]
    r = client.get("/api/ledger/export?month=2026-05&format=xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Date" and len(rows) == 4
    assert rows[1][1] == "'=1+1" and ws.cell(2, 2).data_type == "s"
    assert rows[2][5] == 2500.5 and rows[2][8] == 450.09 and rows[2][9] == 2950.59
    assert isinstance(rows[1][9], int | float) and rows[-1][0] == "TOTAL"
    assert rows[-1][9] == round(1180 + 2950.59, 2)


def test_export_empty_month_and_validation(client: TestClient) -> None:
    setup(client)
    r = client.get("/api/ledger/export?month=2026-05&format=csv")
    assert len(r.content.decode("utf-8-sig").strip().splitlines()) == 2  # header + total
    assert client.get("/api/ledger/export?month=2026-05&format=exe").status_code == 422
    assert client.get("/api/ledger/export?format=csv").status_code == 422
    assert client.get('/api/ledger/export?month="x/../y"').status_code == 422


def test_ledger_tenant_isolation(client: TestClient) -> None:
    biz_a, conn = setup(client, "a@example.com")
    add_bill(conn, biz_a, vendor="A Secret Vendor")  # type: ignore[arg-type]
    add_party(conn, biz_a, "A Secret Vendor")  # type: ignore[arg-type]
    client.post("/api/auth/logout")
    setup(client, "b@example.com")
    assert client.get("/api/ledger").json()["total_count"] == 0
    assert client.get("/api/ledger?q=secret").json()["total_count"] == 0
    assert client.get("/api/ledger/filters").json()["vendors"] == []
    assert client.get("/api/dashboard").json()["bills"] == 0
    csv_body = client.get("/api/ledger/export?month=2026-05&format=csv").content.decode("utf-8-sig")
    assert "Secret" not in csv_body
    assert client.get("/api/ledger").status_code == 200
    client.post("/api/auth/logout")
    assert client.get("/api/ledger").status_code == 401
    assert client.get("/api/ledger/export?month=2026-05").status_code == 401
